# -*- coding: utf-8 -*-
"""
pre_trade.py
진입 '직전'에 돌리는 경고 조회. 사후 기록(mistake_log.py)이 쌓은 과거 실수를 진입 전에 다시 보여준다.

사용법:
    python pre_trade.py SK하이닉스 --side long
    옵션: --days 90 | --xlsx 일지.xlsx | --no-gate (관문 점검 생략)

출력(경고 5줄 + 관문):
    1~3줄  같은 종목(tickers 부분일치) 과거 실수 최근순 3건. 모자라면 빨간 태그와 같은 태그 실수로 채운다.
    4줄    최근 N일(기본 90) 빨간(3회↑) 태그 상위 3개
    5줄    그 태그의 '다음엔(규칙)'
    관문   trading/my 시트에서 이 종목의 '미청산' 행(청산가 빈 행)을 찾아 손절가·손실금액·청산조건이 있는지 본다.
           행이 없거나 하나라도 비어 있으면 '관문 미통과'.

데이터 출처: 매매일지/evernote/entries.jsonl + 엑셀 '반복실수' 시트의 수동 행(mistake_log.read_manual).
관문의 열: 1행 헤더로 찾는다(update_trading_journal.HEADER_ALIASES).
    '손절가'·'손실금액'·'청산조건' 열이 시트에 생기면 그 열을 우선 보고,
    없으면(현재) 같은 행의 '시나리오'·'매매근거' 글에서 숫자가 붙은 표현을 찾는다.
"""
import argparse
import datetime as dt
import json
import os
import re
import shutil
import sys
import tempfile

import openpyxl

import mistake_log as ml
import update_trading_journal as u

RED = 3                     # 최근 N일 안에 이 횟수 이상이면 빨간 태그
# 같은 종목의 다른 표기. 엑셀 시트는 티커(SKHX), 에버노트는 한글이라 부분일치만으로는 못 만난다.
ALIASES = [
    {'SK하이닉스', '하이닉스', 'SKHX', 'SKH', '000660'},
    {'삼성전자', '삼전', '005930'},
    {'비트코인', 'BTC', 'BTCUSDT'},
    {'이더리움', 'ETH', 'ETHUSDT'},
]
# 관문 항목: (이름, '손절가' 같은 전용 열 후보, 글에서 찾는 패턴)
NUM = r'\d[\d,\.]*'
GATE = [
    ('손절가', ('손절가',),
     re.compile(rf'(손절|SL\b|스탑|stop)[^\n]{{0,14}}{NUM}|{NUM}\s*[가-힣$%]{{0,2}}\s*(손절|이탈\s*시|SL\b)', re.I)),
    ('손실금액', ('손실금액', '손실 금액'),
     re.compile(rf'(손실|리스크|risk|최대\s*손실|1R)[^\n]{{0,12}}{NUM}|-\s*{NUM}\s*(%|불|원|달러)', re.I)),
    ('청산조건', ('청산조건', '청산 조건'),
     re.compile(rf'(TP|청산|목표|익절|트레일링|타겟|탈출)[^\n]{{0,14}}{NUM}|{NUM}\s*[가-힣$%]{{0,3}}\s*(청산|익절|TP\b|목표|탈출)', re.I)),
]
SIDE_WORDS = {'long': ('롱', '현물', '매수', 'long'), 'short': ('숏', '매도', 'short')}


def norm(s):
    return re.sub(r'[\s()（）\[\]]', '', str(s or '')).lower()


def keys_for(name):
    """검색어와 같은 종목으로 볼 이름들(정규화)."""
    n = norm(name)
    out = {n}
    for grp in ALIASES:
        g = {norm(x) for x in grp}
        if any(n in x or x in n for x in g if len(x) >= 2):
            out |= g
    return out


def matches(keys, text):
    t = norm(text)
    return bool(t) and any(k and (k in t or (len(t) >= 2 and t in k)) for k in keys)


def open_wb(path, **kw):
    """엑셀이 열려 있어 읽기가 막히면 임시 복사본으로 읽는다."""
    try:
        return openpyxl.load_workbook(path, **kw)
    except PermissionError:
        tmp = os.path.join(tempfile.mkdtemp(), 'copy.xlsx')
        shutil.copy2(path, tmp)
        return openpyxl.load_workbook(tmp, **kw)


def load_rows(wb):
    """에버노트 적재분 + 시트의 수동 행을 한 목록으로. 태그가 있는 것만."""
    rows = []
    if os.path.exists(ml.ENTRIES):
        with open(ml.ENTRIES, encoding='utf-8') as f:
            for line in f:
                e = json.loads(line)
                if 'dup_of' in e or not e['tags']:
                    continue
                rows.append({'date': dt.date.fromisoformat(e['date']), 'tags': e['tags'],
                             'tickers': e['tickers'], 'what': e['title'],
                             'rules': e['rules'], 'result': e.get('outcome') or ''})
    manual = ml.read_manual(wb[ml.SHEET]) if ml.SHEET in wb.sheetnames else []
    for m in manual:
        d = m['날짜']
        rows.append({'date': d.date() if isinstance(d, dt.datetime) else d,
                     'tags': [t.strip() for t in str(m.get('태그') or '').split(',') if t.strip()],
                     'tickers': [x.strip() for x in str(m.get('종목') or '').split(',') if x.strip()],
                     'what': str(m.get('무슨 일이 있었나') or ''),
                     'rules': [m['다음엔(규칙)']] if m.get('다음엔(규칙)') else [],
                     'result': str(m.get('결과·손익') or '')})
    rows.sort(key=lambda r: r['date'], reverse=True)          # 최근순
    return rows


def red_tags(rows, today, days):
    cnt = {}
    for r in rows:
        if (today - r['date']).days <= days:
            for t in r['tags']:
                cnt[t] = cnt.get(t, 0) + 1
    return sorted(((t, n) for t, n in cnt.items() if n >= RED), key=lambda x: (-x[1], x[0]))[:3]


# 규칙 문장이 그 태그 얘기인지 가리는 핵심어. 일기 한 항목에 태그 여러 개·규칙 여러 개가 섞여 있어
# 핵심어 없이 고르면 X1(익절)에 "펀더멘탈보다 가격을 우선한다" 같은 문장이 붙었다(2026-10-19).
TAG_KEYWORDS = {
    'E1': ('돌파', '눌림', '패턴', '완성', '확인'), 'E2': ('충분히', '하락한 후', '떨어'),
    'E3': ('검증', '공부', '맹신', '리포트'), 'E4': ('사이즈', '비중', '복구', '만회'),
    'E5': ('업종', '따라'),
    'E6': ('가격', '매크로', '펀더멘탈', '수급', '추세가 이탈'),
    'X1': ('청산', '익절', '매도', '차익', '수익 실현', '털어', '판다'), 'X2': ('트레일링', '본절'),
    'X3': ('손절',), 'X4': ('재진입', '다시 사', '추세'),
    'R1': ('신용', '레버리지', '담보'), 'R2': ('계약', '사이즈', '랏수', '손절 금액'),
    'R3': ('물타', '추가 매수'), 'R4': ('헤지', '헷지', '헷징'),
    'P1': ('낙관', '장밋빛', '환희', '틀릴', '과신', '반대', '보수적', '하락 계획'),
    'P2': ('FOMO', '조급', '욕구'), 'P3': ('재진입', '당일', '오기'), 'P4': ('방심', '리스크'),
    'P5': ('휴가', '피로', '음주'), 'O1': ('주문', '수량', '방향'), 'O2': ('알림', '알람'),
}


def rule_for(rows, tag, used, width=34):
    """태그의 '다음엔(규칙)'. 그 태그 핵심어가 든 문장을 먼저, 첫 태그가 이 태그인 행을 먼저 본다.
    다른 태그가 이미 쓴 문장은 건너뛴다."""
    cand = [r for r in rows if tag in r['tags'] and r['rules']]
    cand.sort(key=lambda r: (r['tags'][0] != tag, len(r['tags']) > 2))   # 안정 정렬 → 첫 태그가 이 태그인 행 우선, 같으면 최근순
    kws = TAG_KEYWORDS.get(tag, ())
    for need_kw in (True, False):
        for r in cand:
            for rule in r['rules']:
                if need_kw and not any(k in str(rule) for k in kws):
                    continue
                rule = ml.clip(str(rule), width)
                if rule not in used:
                    used.add(rule)
                    return rule
    return next((t[3] for t in ml.TAGS if t[0] == tag), '')


def short(r):
    t = ' '.join(r['tags'])
    res = f" → {ml.clip(r['result'], 18)}" if r['result'] else ''
    return f"{r['date']:%y-%m-%d} {t} {ml.clip(r['what'], 34)}{res}"


def sheet_cols(ws):
    """1행 헤더 → 열 번호. 필요한 열의 이름 후보는 update_trading_journal.HEADER_ALIASES 를 따른다."""
    hdr, _ = u.resolve_columns(ws)
    idx = {name: openpyxl.utils.column_index_from_string(L) for name, L in hdr.items()}
    def find(key):
        return next((idx[a] for a in u.HEADER_ALIASES[key] if a in idx), None)
    return idx, {'sym': find('C_SYM'), 'side': find('C_SIDE'), 'entry': find('C_ENTRY'),
                 'exit': find('C_EXIT'), 'date': find('C_DATE')}


def gate_check(wb, keys, side):
    """trading/my 에서 이 종목의 미청산 행을 찾아 [(시트, 행, 종목, [빠진 항목])] 로."""
    found = []
    for name in ('trading', 'my'):
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        idx, c = sheet_cols(ws)
        if not (c['sym'] and c['entry'] and c['exit']):
            continue
        text_cols = [idx[h] for h in ('시나리오', '매매근거', '매매 근거') if h in idx]
        for r in range(2, ws.max_row + 1):
            sym = ws.cell(r, c['sym']).value
            if not matches(keys, sym) or ws.cell(r, c['exit']).value not in (None, ''):
                continue                                       # 종목 불일치이거나 이미 청산
            if side and c['side']:
                sv = str(ws.cell(r, c['side']).value or '').lower()
                if sv and not any(w in sv for w in SIDE_WORDS[side]):
                    continue
            text = '\n'.join(str(ws.cell(r, t).value or '') for t in text_cols)
            missing = []
            for label, headers, pat in GATE:
                col = next((idx[h] for h in headers if h in idx), None)
                ok = bool(ws.cell(r, col).value) if col else bool(pat.search(text))
                if not ok:
                    missing.append(label)
            found.append((name, r, str(sym), missing))
    return found


def recent_recurrences(rows, today, days=30):
    """최근 days일 안에 난 실수 중, 같은 태그가 직전 days일 안에도 있었던 건 (최근순)."""
    last, out = {}, []
    for r in sorted(rows, key=lambda r: r['date']):
        for t in r['tags']:
            prev = last.get(t)
            if (today - r['date']).days <= days and prev and (r['date'] - prev).days <= days:
                out.append({'date': r['date'], 'tag': t, 'prev': prev,
                            'gap': (r['date'] - prev).days, 'what': r['what']})
            last[t] = r['date']
    return out[::-1]


# 체크리스트_2611 에 새로 넣은 항목이 막는 태그(mistake_log.TAGS 의 '없음 — 추가 후보'를 대신한다)
GUARD_2611 = {
    'X1': '수익 시 트레일링 조건(예: +10% 후 -5%)을 미리 걸었다',
    'X2': '수익 시 트레일링 조건(예: +10% 후 -5%)을 미리 걸었다',
    'X3': '진입과 동시에 손절 주문(TP/SL·자동감시)을 넣었다',
    'R1': '신용 사용액이 계좌의 1/3 이내다',
}


def repeat_report(xlsx=None, days=90, recur_days=30, today=None):
    """텔레그램 반복실수 알림용 자료. 읽을 데이터가 없으면 None — 호출한 쪽 발송을 막지 않는다."""
    today = today or dt.date.today()
    try:
        wb = open_wb(xlsx or u.DEFAULT_XLSX, data_only=True)
    except Exception:
        wb = openpyxl.Workbook()                               # 엑셀을 못 읽으면 에버노트 적재분만으로
    rows = load_rows(wb)
    if not rows:
        return None
    used, red = set(), []
    for t, n in red_tags(rows, today, days):
        tag = next((x for x in ml.TAGS if x[0] == t), (t, t, '', '없음'))
        guard = GUARD_2611.get(t) or ('' if tag[3].startswith('없음') else tag[3])
        red.append({'tag': t, 'name': tag[1], 'desc': tag[2], 'count': n,
                    'last': max(r['date'] for r in rows if t in r['tags']),
                    'rule': rule_for(rows, t, used, width=80), 'guard': guard})
    return {'today': today, 'days': days, 'recur_days': recur_days,
            'n_recent': sum(1 for r in rows if (today - r['date']).days <= days),
            'red': red, 'recur': recent_recurrences(rows, today, recur_days)}


def main():
    ap = argparse.ArgumentParser(description='진입 전 경고 조회')
    ap.add_argument('ticker', help='종목명(부분일치, 예: SK하이닉스)')
    ap.add_argument('--side', choices=('long', 'short'), help='방향(관문 점검의 행 선택에 쓴다)')
    ap.add_argument('--days', type=int, default=90)
    ap.add_argument('--xlsx', default=u.DEFAULT_XLSX)
    ap.add_argument('--no-gate', action='store_true')
    args = ap.parse_args()

    wb = open_wb(args.xlsx, data_only=True)
    rows = load_rows(wb)
    today = dt.date.today()
    keys = keys_for(args.ticker)
    side_txt = {'long': '롱', 'short': '숏'}.get(args.side, '')
    print(f'[{args.ticker} {side_txt}] {today}'.rstrip())

    red = red_tags(rows, today, args.days)
    red_set = {t for t, _ in red}
    same = [r for r in rows if any(matches(keys, t) for t in r['tickers'])]
    picked = same[:3]
    if len(picked) < 3:                                        # 같은 종목이 모자라면 같은 빨간 태그로
        extra = [r for r in rows if r not in picked and red_set & set(r['tags'])]
        picked += extra[:3 - len(picked)]
    for r in picked:
        mark = '종목' if r in same else '태그'
        print(f'  ({mark}) {short(r)}')
    if not picked:
        print('  과거 기록: 이 종목·빨간 태그 실수 없음')

    if red:
        print(f"  빨간 태그({args.days}일): " + ' · '.join(
            f'{t} {ml.TAG_NAME.get(t, "")} {n}회' for t, n in red))
        used = set()
        print('  다음엔: ' + ' │ '.join(f'{t} {rule_for(rows, t, used)}' for t, _ in red))
    else:
        print(f'  빨간 태그({args.days}일): 없음')

    if args.no_gate:
        return
    found = gate_check(wb, keys, args.side)
    if not found:
        print('  [관문 미통과] trading/my 시트에 이 종목의 미청산 행이 없음 — '
              '진입 전에 행을 만들고 손절가·손실금액·청산조건부터 쓰세요')
        return
    for name, r, sym, missing in found:
        if missing:
            print(f"  [관문 미통과] {name} {r}행 {sym}: {', '.join(missing)} 비어 있음")
        else:
            print(f'  [관문 통과] {name} {r}행 {sym}: 손절가·손실금액·청산조건 있음')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
