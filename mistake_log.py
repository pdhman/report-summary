# -*- coding: utf-8 -*-
"""
mistake_log.py
매매일지에 '반복실수' 시트를 만들고, 자주 하는 실수를 계속 기록한다.

사용법:
    python mistake_log.py                       # 시트 갱신(에버노트 적재분 + 손으로/스크립트로 추가한 행 유지)
    python mistake_log.py --add 2026-10-17 --tags R1,X1 --what "신용 물량을 반등에서 못 줄임" \
        --ticker 삼성전자 --account my --result "-120만원" --rule "신용은 계좌 1/3 이내" --source "my 시트 165행"
    옵션: --xlsx 일지.xlsx | --dry-run | --force | --no-backup

시트 구성(반복실수):
    위  — 실수 유형 요약: 태그별 누적·최근 90일·최근 12개월·마지막 발생(전부 수식, 아래 기록에서 자동 계산)
          최근 90일 3회 이상 빨강, 1회 이상 주황.
    아래 — 실수 기록: 한 행 = 실수 한 건(태그는 쉼표로 여러 개). 날짜순.
          '재발' 열 = 같은 태그가 직전에 나온 뒤 며칠 만인지(스크립트가 매번 다시 계산).

데이터 출처:
    매매일지/evernote/entries.jsonl (에버노트 일기·매매일지 적재분, 태그 있는 항목만)
    + 시트에 이미 있는 행 중 출처가 '에버노트'로 시작하지 않는 행(손으로 쓴 행·--add 행)은 그대로 보존한다.
태그 정의: 매매일지/실수태그_v0.md
"""
import argparse
import datetime as dt
import json
import os
import shutil
import sys

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from update_trading_journal import (BACKUP_DIR, DEFAULT_XLSX, KEEP_BACKUPS, center_vertical,
                                    prune_backups)

HERE = os.path.dirname(os.path.abspath(__file__))
ENTRIES = os.path.join(HERE, '매매일지', 'evernote', 'entries.jsonl')
SHEET = '반복실수'
AFTER_SHEET = '마인드'            # 이 시트 바로 뒤에 둔다

# 코드, 이름, 정의, 막는 장치(체크리스트_2610 항목 또는 '없음')
TAGS = [
    ('E1', '조기 진입', '패턴(컵핸들·넥라인·박스 돌파) 완성 전 진입', '셋업 조건 충족 후 진입'),
    ('E2', '조정 초입 성급 매수', '하락 시작 직후 "많이 빠졌다"고 매수', '셋업 조건 충족 후 진입'),
    ('E3', '검증 없는 진입', '추천·내러티브·리포트만 믿고 진입', '타인 의견으로 계획 변경 금지'),
    ('E4', '복구 조급 과대 비중', '손실 복구 조급함으로 비중 확대', '전일 손익이 사이즈·심리에 영향 없는지'),
    ('E5', '연쇄 희망 매수', '"같은 업종이니 따라 오르겠지"', '셋업 조건 충족 후 진입'),
    ('X1', '익절 실기 / 엑시트 플랜 부재', '목표가·매도 자리 없이 보유하다 수익 반납', '없음 — 추가 후보'),
    ('X2', '트레일링·본절 미설정', '수익 후 손절선을 진입가 위로 올리지 않음', '없음 — 추가 후보'),
    ('X3', '손절 지연·미이행', '지지 이탈·아이디어 훼손에도 보유', '진입 전 손절가·손실금액(R) 확정'),
    ('X4', '조기 매도 후 재진입 실패', '돌파를 추세로 못 보고 일찍 팔고 못 탐', '없음'),
    ('R1', '신용·레버리지 과용', '신용 담보 위험·레버리지 ETF 조기 매수·마진콜', '없음 — 추가 후보'),
    ('R2', '포지션 규모 > 손절 상한', '손절가 역산 없이 원하는 크기로 진입', '일일 손실 한도를 숫자로 적었다'),
    ('R3', '물타기', '손실 포지션에 추가 매수', '분할매수 계획대로 진입'),
    ('R4', '헤지 선택 오류', '헤지로 산 자산이 같이 빠짐', '없음'),
    ('P1', '긍정 편향·환희', '상승 뉴스만 보고 부정 신호 무시', '조급함/탐욕 40 이상이면 신규 진입 금지'),
    ('P2', 'FOMO', '시장이 안 빠져서 매수 욕구를 못 참음', '조급함/탐욕 40 이상이면 신규 진입 금지'),
    ('P3', '오기·재진입', '손절 후 당일 재진입, 연속 손절 뒤 포지션 키움', '손절 종목 당일 재진입 금지'),
    ('P4', '방심', '연속 반등 뒤 리스크 감각 둔화', '없음'),
    ('P5', '환경 요인 대응 부재', '휴가·피로·음주 중 방치', '없음'),
    ('O1', '주문 실수', '매수/매도 방향·수량 오입력', '주문 전 방향·수량 재확인'),
    ('O2', '알림만 설정, 행동 규칙 없음', '알람이 울려도 "더 갈 것"이라 판단', '없음'),
]
TAG_NAME = {t[0]: t[1] for t in TAGS}
ACCOUNT = {'prop': 'prop', 'my': '국내(my)', 'fx-company': 'FX(회사)', 'personal': '개인', 'unknown': ''}

COLS = ['날짜', '태그', '실수 유형', '종목', '계좌', '무슨 일이 있었나', '결과·손익', '다음엔(규칙)', '출처', '재발(일)', 'ID']
WIDTHS = [11, 10, 24, 18, 9, 58, 22, 44, 26, 8, 22]
SUM_ROW0 = 4                       # 요약 머리글 행
LOG_ROW0 = SUM_ROW0 + len(TAGS) + 4  # 기록 머리글 행

FILL_TITLE = PatternFill('solid', fgColor='1F3864')
FILL_HEAD = PatternFill('solid', fgColor='C00000')
FILL_SUB = PatternFill('solid', fgColor='FCE4D6')
FILL_RED = PatternFill('solid', start_color='FF9999', end_color='FF9999')      # 조건부서식(dxf)은 bgColor를 써서 둘 다 지정
FILL_ORANGE = PatternFill('solid', start_color='FFE0B2', end_color='FFE0B2')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='center')
CENTER = Alignment(horizontal='center', vertical='center')


def clip(s, n):
    s = ' '.join((s or '').split())
    return s if len(s) <= n else s[:n - 1] + '…'


def from_entries():
    """에버노트 적재분 중 실수 태그가 있는 항목을 기록 행으로 바꾼다."""
    rows = []
    if not os.path.exists(ENTRIES):
        print(f'에버노트 적재 파일 없음: {ENTRIES} (손으로 쓴 행만 유지)')
        return rows
    with open(ENTRIES, encoding='utf-8') as f:
        for line in f:
            e = json.loads(line)
            if 'dup_of' in e or not e['tags']:
                continue
            mistakes = [l['text'] for l in e.get('lines') or [] if l['kind'] == 'mistake']
            body = ' / '.join(mistakes) if mistakes else e['text'].split('\n', 1)[-1]
            rows.append({
                '날짜': dt.date.fromisoformat(e['date']),
                '태그': ', '.join(e['tags']),
                '종목': ', '.join(e['tickers'][:3]),
                '계좌': ACCOUNT.get(e['account'], e['account']),
                '무슨 일이 있었나': clip(f"{e['title']} — {body}", 160),
                '결과·손익': clip(e.get('outcome') or '', 60),
                '다음엔(규칙)': clip(' / '.join(e['rules'][:2]), 120),
                '출처': f"에버노트 · {e['source']['note_title']}",
                'ID': e['id'],
            })
    return rows


def read_manual(ws):
    """기존 시트의 기록 중 에버노트가 아닌 행(손으로 쓴 행·--add 행)을 읽는다."""
    rows = []
    if ws is None:
        return rows
    head = {ws.cell(LOG_ROW0, c).value: c for c in range(1, ws.max_column + 1)}
    if '날짜' not in head:
        return rows
    for r in range(LOG_ROW0 + 1, ws.max_row + 1):
        src = ws.cell(r, head['출처']).value or ''
        d = ws.cell(r, head['날짜']).value
        if d is None or str(src).startswith('에버노트'):
            continue
        row = {k: ws.cell(r, c).value for k, c in head.items() if k in COLS}
        if isinstance(d, dt.datetime):
            row['날짜'] = d.date()
        rows.append(row)
    return rows


def add_recurrence(rows):
    """재발(일): 이 행의 태그 중 하나라도 직전에 나온 날로부터 며칠 만인지(가장 짧은 값)."""
    last = {}
    for r in rows:
        gaps = []
        for t in [t.strip() for t in str(r['태그']).split(',') if t.strip()]:
            if t in last:
                gaps.append((r['날짜'] - last[t]).days)
            last[t] = r['날짜']
        r['재발(일)'] = min(gaps) if gaps else None


def write_sheet(wb, rows):
    old_idx = wb.sheetnames.index(SHEET) if SHEET in wb.sheetnames else None
    if old_idx is not None:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    pos = old_idx if old_idx is not None else (
        wb.sheetnames.index(AFTER_SHEET) + 1 if AFTER_SHEET in wb.sheetnames else len(wb.sheetnames) - 1)
    wb.move_sheet(ws, offset=pos - wb.sheetnames.index(SHEET))
    ws.sheet_properties.tabColor = 'C00000'

    ws['A1'] = '반복실수 — 같은 실수를 세고, 다시 나오면 바로 보이게'
    ws['A1'].font = Font(size=14, bold=True, color='FFFFFF')
    ws['A1'].fill = FILL_TITLE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLS))
    ws['A2'] = ('아래 "실수 기록"에 한 줄씩 추가하면 요약이 자동 계산됩니다. 태그는 쉼표로 여러 개(예: R1, X1). '
                '정의: 매매일지/실수태그_v0.md · 최근 90일 3회↑ 빨강 / 1회↑ 주황')
    ws['A2'].font = Font(size=9, color='808080')
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLS))

    # ---- 요약 ----
    first, last = LOG_ROW0 + 1, LOG_ROW0 + max(len(rows), 1) + 500   # 앞으로 500행 추가분까지 범위에 넣는다
    tag_rng = f'$B${first}:$B${last}'
    date_rng = f'$A${first}:$A${last}'
    heads = ['코드', '실수 유형', '정의', '누적', '최근 90일', '최근 12개월', '마지막 발생', '막는 장치(체크리스트)']
    for c, h in enumerate(heads, 1):
        cell = ws.cell(SUM_ROW0, c, h)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = FILL_HEAD
        cell.alignment = CENTER
        cell.border = BOX
    ws.cell(SUM_ROW0 - 1, 1, '실수 유형 요약').font = Font(size=12, bold=True)
    count = {t[0]: sum(1 for r in rows if t[0] in str(r['태그'])) for t in TAGS}
    order = sorted(TAGS, key=lambda t: -count[t[0]])            # 많이 한 실수부터
    for i, (code, name, desc, guard) in enumerate(order):
        r = SUM_ROW0 + 1 + i
        pat = f'"*{code}*"'
        vals = [code, name, desc,
                f'=COUNTIF({tag_rng},{pat})',
                f'=COUNTIFS({tag_rng},{pat},{date_rng},">="&(TODAY()-90))',
                f'=COUNTIFS({tag_rng},{pat},{date_rng},">="&EDATE(TODAY(),-12))',
                f'=IFERROR(IF(_xlfn.MAXIFS({date_rng},{tag_rng},{pat})=0,"",_xlfn.MAXIFS({date_rng},{tag_rng},{pat})),"")',
                guard]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.border = BOX
            cell.alignment = CENTER if c in (1, 4, 5, 6, 7) else WRAP
        ws.cell(r, 1).font = Font(bold=True)
        ws.cell(r, 7).number_format = 'yyyy-mm-dd'
        if guard.startswith('없음'):
            ws.cell(r, 8).font = Font(color='C00000', bold=guard.endswith('후보'))
    s0, s1 = SUM_ROW0 + 1, SUM_ROW0 + len(TAGS)
    ws.conditional_formatting.add(f'E{s0}:E{s1}', CellIsRule(operator='greaterThanOrEqual', formula=['3'], fill=FILL_RED))
    ws.conditional_formatting.add(f'E{s0}:E{s1}', CellIsRule(operator='greaterThanOrEqual', formula=['1'], fill=FILL_ORANGE))

    # ---- 기록 ----
    ws.cell(LOG_ROW0 - 1, 1, '실수 기록').font = Font(size=12, bold=True)
    for c, h in enumerate(COLS, 1):
        cell = ws.cell(LOG_ROW0, c, h)
        cell.font = Font(bold=True)
        cell.fill = FILL_SUB
        cell.alignment = CENTER
        cell.border = BOX
    for i, row in enumerate(rows):
        r = LOG_ROW0 + 1 + i
        row['실수 유형'] = ', '.join(TAG_NAME.get(t.strip(), t.strip()) for t in str(row['태그']).split(','))
        for c, key in enumerate(COLS, 1):
            cell = ws.cell(r, c, row.get(key))
            cell.border = BOX
            cell.alignment = CENTER if key in ('날짜', '태그', '계좌', '재발(일)') else WRAP
        ws.cell(r, 1).number_format = 'yyyy-mm-dd'
        ws.cell(r, 11).font = Font(size=8, color='A6A6A6')
        gap = row.get('재발(일)')
        if gap is not None and gap <= 30:
            ws.cell(r, 10).fill = FILL_RED
    for c, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ws.column_dimensions['C'].width = max(ws.column_dimensions['C'].width, 24)
    ws.auto_filter.ref = f'A{LOG_ROW0}:{openpyxl.utils.get_column_letter(len(COLS))}{LOG_ROW0 + max(len(rows), 1)}'
    center_vertical(ws)
    return ws


def main():
    ap = argparse.ArgumentParser(description='반복실수 시트 갱신·기록 추가')
    ap.add_argument('--xlsx', default=DEFAULT_XLSX)
    ap.add_argument('--add', metavar='YYYY-MM-DD', help='새 실수 날짜')
    ap.add_argument('--tags', help='태그(쉼표), 예: R1,X1')
    ap.add_argument('--what', help='무슨 일이 있었나')
    ap.add_argument('--ticker', default='')
    ap.add_argument('--account', default='')
    ap.add_argument('--result', default='')
    ap.add_argument('--rule', default='')
    ap.add_argument('--source', default='수동')
    ap.add_argument('--dry-run', action='store_true')
    ap.add_argument('--force', action='store_true')
    ap.add_argument('--no-backup', action='store_true')
    args = ap.parse_args()

    if args.add:
        if not (args.tags and args.what):
            sys.exit('--add 에는 --tags 와 --what 이 필요합니다.')
        bad = [t for t in args.tags.replace(' ', '').split(',') if t not in TAG_NAME]
        if bad:
            sys.exit(f'모르는 태그: {bad} (가능: {", ".join(TAG_NAME)})')
        if args.source.startswith('에버노트'):
            sys.exit('--source 는 "에버노트"로 시작하면 안 됩니다(갱신 때 덮어써짐).')

    lock = os.path.join(os.path.dirname(os.path.abspath(args.xlsx)), '~$' + os.path.basename(args.xlsx))
    if os.path.exists(lock) and not (args.dry_run or args.force):
        sys.exit(f'엑셀에서 열려 있는 것 같습니다({os.path.basename(lock)}). 닫고 다시 실행하세요.')

    wb = openpyxl.load_workbook(args.xlsx)
    manual = read_manual(wb[SHEET] if SHEET in wb.sheetnames else None)
    if args.add:
        new = {'날짜': dt.date.fromisoformat(args.add), '태그': ', '.join(args.tags.replace(' ', '').split(',')),
               '종목': args.ticker, '계좌': args.account, '무슨 일이 있었나': args.what,
               '결과·손익': args.result, '다음엔(규칙)': args.rule, '출처': args.source,
               'ID': f'M-{args.add}-{len(manual) + 1:02d}'}
        dup = any(m['날짜'] == new['날짜'] and m['태그'] == new['태그'] and m['무슨 일이 있었나'] == new['무슨 일이 있었나']
                  for m in manual)
        if dup:
            print('같은 날짜·태그·내용의 행이 이미 있어 추가하지 않습니다.')
        else:
            manual.append(new)
    rows = from_entries() + manual
    rows.sort(key=lambda r: (r['날짜'], str(r.get('ID') or '')))
    add_recurrence(rows)
    write_sheet(wb, rows)

    cnt = {}
    for r in rows:
        for t in str(r['태그']).split(','):
            cnt[t.strip()] = cnt.get(t.strip(), 0) + 1
    top = sorted(cnt.items(), key=lambda x: -x[1])[:5]
    print(f'기록 {len(rows)}행 (에버노트 {len(rows) - len(manual)} · 수동 {len(manual)}) · 상위 태그 {top}')
    recent = [r for r in rows if r.get('재발(일)') is not None and r['재발(일)'] <= 30][-3:]
    for r in recent:
        print(f"  30일 내 재발: {r['날짜']} {r['태그']} ({r['재발(일)']}일)")

    if args.dry_run:
        print('--dry-run: 저장하지 않음')
        return
    if not args.no_backup:
        os.makedirs(BACKUP_DIR, exist_ok=True)
        bak = os.path.join(BACKUP_DIR, f'{os.path.basename(args.xlsx)}.{dt.datetime.now():%Y%m%d_%H%M%S}.bak')
        shutil.copy2(args.xlsx, bak)
        print(f'백업: {bak}')
        prune_backups(BACKUP_DIR, os.path.basename(args.xlsx), KEEP_BACKUPS)
    wb.save(args.xlsx)
    print(f'저장: {args.xlsx} [{SHEET}]')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
